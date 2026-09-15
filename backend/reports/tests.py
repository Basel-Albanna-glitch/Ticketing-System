import io
from datetime import timedelta

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from accounts.models import CustomerLicense, User
from tickets.models import Category, Ticket


class ReportsTests(APITestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.admin = User.objects.create_user(
            username='boss', full_name='The Admin', role=User.Role.ADMIN
        )
        self.agent = User.objects.create_user(
            username='ag', full_name='An Agent', role=User.Role.AGENT
        )
        self.customer = User.objects.create_user(
            username='cust', full_name='Zaytoun Grill', role=User.Role.CUSTOMER,
            software_types=['POS'],
        )
        self.quiet_customer = User.objects.create_user(
            username='quiet', full_name='No Tickets Ltd', role=User.Role.CUSTOMER
        )
        self.category = Category.objects.create(name='Hardware')
        self.client.force_authenticate(self.admin)

    def _ticket(self, status='open', rating=None, resolved=False):
        ticket = Ticket.objects.create(
            subject='Till offline', customer=self.customer, assigned_agent=self.agent,
            category=self.category, status=status, rating=rating,
        )
        if resolved:
            ticket.resolved_at = ticket.created_at + timedelta(hours=4)
            ticket.save()
        return ticket

    # --- customer activity ---

    def test_customer_activity_counts_and_excludes_quiet_customers(self):
        self._ticket()
        self._ticket(status='closed', rating=4, resolved=True)

        response = self.client.get('/api/reports/customers/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1, 'customers with no tickets are dropped')

        row = response.data[0]
        self.assertEqual(row['customer_name'], 'Zaytoun Grill')
        self.assertEqual(row['software_type'], 'POS')
        self.assertEqual(row['ticket_count'], 2)
        self.assertEqual(row['open_count'], 1)
        self.assertEqual(row['closed_count'], 1)
        self.assertEqual(row['avg_rating'], 4.0)
        self.assertEqual(row['avg_resolution_hours'], 4.0)

    def test_customer_activity_respects_the_date_range(self):
        self._ticket()
        future = (self.today + timedelta(days=5)).isoformat()
        response = self.client.get(
            '/api/reports/customers/', {'date_from': future, 'date_to': future}
        )
        self.assertEqual(response.data, [])

    # --- licenses ---

    def test_license_report_buckets_by_end_date(self):
        for days, name in [(-5, 'Lapsed'), (10, 'Renew soon'), (200, 'Comfortable')]:
            CustomerLicense.objects.create(
                customer=self.customer, name=name,
                end_date=self.today + timedelta(days=days),
            )
        # No end date at all: nothing to report on.
        CustomerLicense.objects.create(customer=self.customer, name='Perpetual')

        response = self.client.get('/api/reports/licenses/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['expired_count'], 1)
        self.assertEqual(response.data['expiring_count'], 1)
        self.assertEqual(response.data['active_count'], 1)

        rows = response.data['results']
        self.assertEqual([r['license_name'] for r in rows],
                         ['Lapsed', 'Renew soon', 'Comfortable'],
                         'soonest end date first')
        self.assertEqual([r['state'] for r in rows], ['expired', 'expiring', 'active'])
        self.assertEqual([r['days_left'] for r in rows], [-5, 10, 200])
        self.assertEqual(rows[0]['customer_id'], self.customer.id)

    def test_license_report_ignores_the_date_range(self):
        """It answers "what needs renewing now", so a narrow window must not hide rows."""
        CustomerLicense.objects.create(
            customer=self.customer, name='Renew soon',
            end_date=self.today + timedelta(days=10),
        )
        past = (self.today - timedelta(days=365)).isoformat()
        response = self.client.get(
            '/api/reports/licenses/', {'date_from': past, 'date_to': past}
        )
        self.assertEqual(len(response.data['results']), 1)

    # --- export ---

    def test_export_returns_a_workbook_with_every_sheet(self):
        self._ticket(status='closed', rating=5, resolved=True)
        CustomerLicense.objects.create(
            customer=self.customer, name='POS licence',
            end_date=self.today + timedelta(days=10),
        )

        response = self.client.get('/api/reports/export/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertIn('attachment; filename="report_', response['Content-Disposition'])

        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            ['Summary', 'By status', 'By priority', 'By category', 'By rating',
             'Agents', 'Customers', 'Licenses'],
        )
        # A fresh Workbook ships with an empty default sheet; it must have been claimed.
        self.assertNotIn('Sheet', workbook.sheetnames)

        summary = dict(workbook['Summary'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary['Total tickets'], 1)
        self.assertEqual(summary['Resolved tickets'], 1)

        licenses = list(workbook['Licenses'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(licenses[0][0], 'Zaytoun Grill')
        self.assertEqual(licenses[0][5], 'Expiring soon')

        customers = list(workbook['Customers'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(customers[0][0], 'Zaytoun Grill')
        self.assertEqual(customers[0][2], 1)

    def test_export_of_an_empty_range_still_produces_a_workbook(self):
        response = self.client.get('/api/reports/export/')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(workbook.sheetnames), 8)

    # --- permissions ---

    def test_new_reports_are_admin_only(self):
        for user in (self.agent, self.customer):
            self.client.force_authenticate(user)
            for url in ('/api/reports/customers/', '/api/reports/licenses/',
                        '/api/reports/export/'):
                self.assertEqual(
                    self.client.get(url).status_code, 403, f'{url} as {user.role}'
                )
