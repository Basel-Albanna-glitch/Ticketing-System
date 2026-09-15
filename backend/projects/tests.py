from datetime import timedelta
from io import BytesIO
from unittest import mock

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from notifications.models import Notification

from .models import Project, Task, TodoFolder, TodoItem, TodoReminder
from .reminders import sweep_todo_reminders


class ClosedProjectTests(TestCase):
    """Closing a project archives it: agents can no longer touch it, and nobody
    can delete it until it is reopened."""

    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='padmin', full_name='Admin', password='testpass123', role=User.Role.ADMIN
        )
        self.agent = User.objects.create_user(
            username='pagent', full_name='Agent', password='testpass123', role=User.Role.AGENT
        )
        self.project = Project.objects.create(name='Delivery', created_by=self.admin)
        self.project.assignees.add(self.agent)

    def close(self):
        Project.objects.filter(pk=self.project.pk).update(status=Project.Status.CLOSED)

    def test_agent_can_edit_an_open_project_they_are_on(self):
        self.client.force_authenticate(user=self.agent)
        response = self.client.patch(
            f'/api/projects/{self.project.pk}/', {'name': 'Renamed'}, format='json'
        )
        self.assertEqual(response.status_code, 200)

    def test_agent_cannot_edit_a_closed_project(self):
        self.close()
        self.client.force_authenticate(user=self.agent)
        response = self.client.patch(
            f'/api/projects/{self.project.pk}/', {'name': 'Renamed'}, format='json'
        )
        self.assertEqual(response.status_code, 403)

    def test_agent_cannot_reopen_a_closed_project(self):
        """Otherwise the freeze would be one request away from being undone."""
        self.close()
        self.client.force_authenticate(user=self.agent)
        response = self.client.patch(
            f'/api/projects/{self.project.pk}/', {'status': 'open'}, format='json'
        )
        self.assertEqual(response.status_code, 403)

    def test_admin_can_still_edit_and_reopen_a_closed_project(self):
        self.close()
        self.client.force_authenticate(user=self.admin)
        response = self.client.patch(
            f'/api/projects/{self.project.pk}/', {'status': 'open'}, format='json'
        )
        self.assertEqual(response.status_code, 200)

    def test_a_closed_project_cannot_be_deleted_even_by_an_admin(self):
        self.close()
        self.client.force_authenticate(user=self.admin)
        response = self.client.delete(f'/api/projects/{self.project.pk}/')
        self.assertEqual(response.status_code, 400)
        self.assertTrue(Project.objects.filter(pk=self.project.pk).exists())

    def test_reopening_makes_it_deletable_again(self):
        """The rule is a speed bump, not a dead end."""
        self.close()
        self.client.force_authenticate(user=self.admin)
        self.client.patch(f'/api/projects/{self.project.pk}/', {'status': 'open'}, format='json')
        self.assertEqual(self.client.delete(f'/api/projects/{self.project.pk}/').status_code, 204)

    def test_can_delete_flag_matches_the_rule(self):
        """The UI hides the delete button on this flag, so it must not promise a
        delete the API will refuse."""
        self.client.force_authenticate(user=self.admin)
        open_payload = self.client.get(f'/api/projects/{self.project.pk}/').json()
        self.assertTrue(open_payload['can_delete'])

        self.close()
        closed_payload = self.client.get(f'/api/projects/{self.project.pk}/').json()
        self.assertFalse(closed_payload['can_delete'])
        # Still editable by an admin — the two flags are deliberately different.
        self.assertTrue(closed_payload['can_edit'])

    def test_a_project_with_unfinished_tasks_cannot_be_closed(self):
        Task.objects.create(
            project=self.project,
            title='Outstanding',
            status=Task.Status.TODO,
            created_by=self.admin,
        )
        self.client.force_authenticate(user=self.admin)
        response = self.client.patch(
            f'/api/projects/{self.project.pk}/', {'status': 'closed'}, format='json'
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Outstanding', str(response.json()))


class MinimalTaskCreationTests(TestCase):
    """A task needs only a project and a title.

    An assignee and a due date were briefly required, which broke both clients:
    the web board relies on unassigned tasks to represent unclaimed work, and the
    mobile board creates standard steps from one-tap template chips that send a
    title and nothing else. These pin the payloads those flows actually send.
    """

    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='tadmin', full_name='Admin', password='testpass123', role=User.Role.ADMIN
        )
        self.project = Project.objects.create(name='Work', created_by=self.admin)
        self.client.force_authenticate(user=self.admin)

    def post(self, **extra):
        payload = {'project': self.project.pk, 'title': 'A task', **extra}
        return self.client.post('/api/tasks/', payload, format='json')

    def test_a_title_alone_is_enough(self):
        self.assertEqual(self.post().status_code, 201)

    def test_the_mobile_template_chip_payload_is_accepted(self):
        """Exactly what mobile/lib/data/misc_repository.dart createTask sends."""
        response = self.client.post(
            '/api/tasks/',
            {
                'project': self.project.pk,
                'title': 'From phone',
                'description': '',
                'status': 'todo',
                'priority': 'medium',
                'assignee_ids': [],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201)

    def test_a_fully_specified_task_is_still_accepted(self):
        response = self.post(assignee_ids=[self.admin.pk], due_date='2026-09-01')
        self.assertEqual(response.status_code, 201)

    def test_date_order_is_still_enforced(self):
        """Dropping the requirement must not drop the checks that remain."""
        response = self.post(start_date='2026-09-10', due_date='2026-09-01')
        self.assertEqual(response.status_code, 400)
        self.assertIn('start_date', response.json())


class TodoReminderTests(TestCase):
    """A to-do can carry several reminders, each either an exact moment or an amount
    before the due date."""

    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='radmin', full_name='Admin', password='testpass123', role=User.Role.ADMIN
        )
        self.client.force_authenticate(user=self.admin)
        self.due = (timezone.now() + timedelta(days=10)).replace(microsecond=0)
        self.moment = (timezone.now() + timedelta(days=2)).replace(microsecond=0)

    def create(self, **payload):
        return self.client.post(
            '/api/todos/', {'title': 'Renew certificate', **payload}, format='json'
        )

    def create_both_kinds(self, offset=1440):
        response = self.create(
            due_at=self.due.isoformat(),
            reminders=[{'offset_minutes': offset}, {'remind_at': self.moment.isoformat()}],
        )
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()['id']

    def test_a_todo_takes_several_reminders_of_either_kind(self):
        todo_id = self.create_both_kinds()
        moments = sorted(TodoReminder.objects.filter(todo_id=todo_id).values_list('remind_at', flat=True))
        self.assertEqual(moments, sorted([self.due - timedelta(days=1), self.moment]))
        self.assertEqual(len(self.client.get(f'/api/todos/{todo_id}/').json()['reminders']), 2)

    def test_a_reminder_before_the_due_date_needs_one(self):
        self.assertEqual(self.create(reminders=[{'offset_minutes': 60}]).status_code, 400)

    def test_an_exact_reminder_needs_no_due_date(self):
        response = self.create(reminders=[{'remind_at': self.moment.isoformat()}])
        self.assertEqual(response.status_code, 201)

    def test_a_reminder_needs_a_time_of_some_kind(self):
        response = self.create(due_at=self.due.isoformat(), reminders=[{}])
        self.assertEqual(response.status_code, 400)

    def test_moving_the_due_date_carries_relative_reminders_and_rearms_them(self):
        todo_id = self.create_both_kinds(offset=60)
        TodoReminder.objects.filter(todo_id=todo_id).update(sent_at=timezone.now())
        later = self.due + timedelta(days=3)
        self.client.patch(f'/api/todos/{todo_id}/', {'due_at': later.isoformat()}, format='json')

        relative = TodoReminder.objects.get(todo_id=todo_id, offset_minutes=60)
        self.assertEqual(relative.remind_at, later - timedelta(hours=1))
        self.assertIsNone(relative.sent_at)
        exact = TodoReminder.objects.get(todo_id=todo_id, offset_minutes__isnull=True)
        self.assertEqual(exact.remind_at, self.moment)
        self.assertIsNotNone(exact.sent_at, 'a fixed reminder does not move with the due date')

    def test_resaving_the_same_reminders_does_not_send_them_again(self):
        todo_id = self.create_both_kinds()
        TodoReminder.objects.filter(todo_id=todo_id).update(sent_at=timezone.now())
        self.client.patch(
            f'/api/todos/{todo_id}/',
            {
                'title': 'Renew the certificate',
                'reminders': [{'offset_minutes': 1440}, {'remind_at': self.moment.isoformat()}],
            },
            format='json',
        )
        self.assertEqual(TodoReminder.objects.filter(todo_id=todo_id).count(), 2)
        self.assertFalse(TodoReminder.objects.filter(todo_id=todo_id, sent_at__isnull=True).exists())

    def test_removing_the_due_date_drops_only_relative_reminders(self):
        todo_id = self.create_both_kinds(offset=60)
        self.client.patch(f'/api/todos/{todo_id}/', {'due_at': None}, format='json')
        kinds = list(TodoReminder.objects.filter(todo_id=todo_id).values_list('offset_minutes', flat=True))
        self.assertEqual(kinds, [None])

    @mock.patch('projects.reminders.email_users')
    def test_reminders_due_together_go_out_as_one_alert(self, _email_users):
        todo = TodoItem.objects.create(title='Order hardware', created_by=self.admin)
        past = timezone.now() - timedelta(minutes=5)
        TodoReminder.objects.create(todo=todo, remind_at=past)
        TodoReminder.objects.create(todo=todo, remind_at=past - timedelta(minutes=1))
        TodoReminder.objects.create(todo=todo, remind_at=timezone.now() + timedelta(days=1))

        self.assertEqual(sweep_todo_reminders(), 1)
        self.assertEqual(Notification.objects.filter(recipient=self.admin).count(), 1)
        self.assertEqual(TodoReminder.objects.filter(todo=todo, sent_at__isnull=False).count(), 2)
        self.assertEqual(sweep_todo_reminders(), 0, 'each reminder goes out once')


class TodoExportLatenessTests(TestCase):
    """A to-do finished after its due time stays marked as late in the export."""

    def test_finishing_after_the_due_time_is_marked_done_late(self):
        client = APIClient()
        admin = User.objects.create_user(
            username='xadmin', full_name='Admin', password='testpass123', role=User.Role.ADMIN
        )
        client.force_authenticate(user=admin)
        due = timezone.now() - timedelta(days=2)
        TodoItem.objects.create(
            title='Late one', created_by=admin, due_at=due, done=True,
            completed_at=due + timedelta(hours=5),
        )
        TodoItem.objects.create(
            title='On time', created_by=admin, due_at=due, done=True,
            completed_at=due - timedelta(hours=1),
        )

        response = client.get('/api/todos/export/')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        statuses = {
            row[0]: row[1] for row in workbook['To-dos'].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(statuses['Late one'], 'Done late')
        self.assertEqual(statuses['On time'], 'Done')
        summary = dict(workbook['Summary'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary['Done late'], 1)


class ShareTodoByDragTests(TestCase):
    """Dragging a private to-do onto the shared list saves is_private and folder_id together.
    Its author can make that move, out of a private folder included; nobody else can."""

    def setUp(self):
        self.client = APIClient()
        self.author = User.objects.create_user(
            username='sauthor', full_name='Author', password='testpass123', role=User.Role.AGENT
        )
        self.colleague = User.objects.create_user(
            username='scolleague', full_name='Colleague', password='testpass123',
            role=User.Role.ADMIN,
        )
        private_folder = TodoFolder.objects.create(
            name='Mine', is_private=True, created_by=self.author
        )
        self.shared_folder = TodoFolder.objects.create(
            name='Team', is_private=False, created_by=self.colleague
        )
        self.todo = TodoItem.objects.create(
            title='Draft plan', created_by=self.author, is_private=True, folder=private_folder
        )

    def share(self, user, folder_id):
        self.client.force_authenticate(user=user)
        return self.client.patch(
            f'/api/todos/{self.todo.pk}/',
            {'is_private': False, 'folder_id': folder_id},
            format='json',
        )

    def test_the_author_can_share_it_unfiled(self):
        self.assertEqual(self.share(self.author, None).status_code, 200)
        self.todo.refresh_from_db()
        self.assertFalse(self.todo.is_private)
        self.assertIsNone(self.todo.folder_id)

    def test_the_author_can_share_it_into_a_shared_folder(self):
        self.assertEqual(self.share(self.author, self.shared_folder.pk).status_code, 200)
        self.todo.refresh_from_db()
        self.assertFalse(self.todo.is_private)
        self.assertEqual(self.todo.folder_id, self.shared_folder.pk)

    def test_nobody_else_can_share_it(self):
        # Someone else's private item is out of reach altogether, admins included.
        self.assertIn(self.share(self.colleague, None).status_code, (400, 403, 404))
        self.todo.refresh_from_db()
        self.assertTrue(self.todo.is_private)


class TodoAssignmentAlertTests(TestCase):
    """Someone a to-do is assigned to gets an alert for it — once, and not for assigning
    themselves."""

    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='aadmin', full_name='Admin', password='testpass123', role=User.Role.ADMIN
        )
        self.ann = User.objects.create_user(
            username='aann', full_name='Ann', password='testpass123', role=User.Role.AGENT
        )
        self.bob = User.objects.create_user(
            username='abob', full_name='Bob', password='testpass123', role=User.Role.AGENT
        )
        self.client.force_authenticate(user=self.admin)

    def alerts_for(self, user):
        return Notification.objects.filter(recipient=user, kind=Notification.Kind.TODO_ASSIGNED)

    def create(self, title, assignees):
        response = self.client.post(
            '/api/todos/', {'title': title, 'assignee_ids': [u.pk for u in assignees]},
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()['id']

    def test_assignees_are_alerted_when_the_todo_is_created(self):
        todo_id = self.create('Order hardware', [self.ann, self.admin])
        self.assertEqual(self.alerts_for(self.ann).get().todo_id, todo_id)
        self.assertFalse(self.alerts_for(self.admin).exists(), 'assigning yourself raises none')

    def test_only_people_newly_added_are_alerted(self):
        todo_id = self.create('Chase supplier', [self.ann])
        self.client.patch(
            f'/api/todos/{todo_id}/', {'assignee_ids': [self.ann.pk, self.bob.pk]}, format='json'
        )
        self.assertEqual(self.alerts_for(self.ann).count(), 1)
        self.assertEqual(self.alerts_for(self.bob).count(), 1)
        # An edit that leaves the assignees alone alerts nobody.
        self.client.patch(f'/api/todos/{todo_id}/', {'title': 'Chase the supplier'}, format='json')
        self.assertEqual(self.alerts_for(self.bob).count(), 1)


class TodoAssigneeCompletionTests(TestCase):
    """A to-do carried by several people tracks each person's share separately.

    `done` remains the state of the whole job; the per-assignee ticks are what it is
    derived from. These pin the invariant that keeps the two from disagreeing: a done
    item has every assignee ticked, an open one does not.
    """

    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='cadmin', full_name='Admin', password='testpass123', role=User.Role.ADMIN
        )
        self.ann = User.objects.create_user(
            username='cann', full_name='Ann', password='testpass123', role=User.Role.AGENT
        )
        self.bob = User.objects.create_user(
            username='cbob', full_name='Bob', password='testpass123', role=User.Role.AGENT
        )
        self.todo = TodoItem.objects.create(title='Two-hander', created_by=self.admin)
        self.todo.assignees.set([self.ann, self.bob])

    def mark(self, user_id, done=True):
        return self.client.post(
            f'/api/todos/{self.todo.pk}/assignees/{user_id}/done/', {'done': done}, format='json'
        )

    def test_the_last_share_closes_the_whole_todo(self):
        self.client.force_authenticate(user=self.ann)
        self.assertEqual(self.mark(self.ann.pk).status_code, 200)
        self.todo.refresh_from_db()
        self.assertFalse(self.todo.done, 'one of two finished is not the job finished')

        self.client.force_authenticate(user=self.bob)
        response = self.mark(self.bob.pk)
        self.todo.refresh_from_db()
        self.assertTrue(self.todo.done)
        self.assertIsNotNone(self.todo.completed_at)
        self.assertEqual(len(response.json()['assignee_completions']), 2)

    def test_taking_a_share_back_reopens_it(self):
        self.client.force_authenticate(user=self.ann)
        self.mark(self.ann.pk)
        self.client.force_authenticate(user=self.bob)
        self.mark(self.bob.pk)
        self.mark(self.bob.pk, done=False)
        self.todo.refresh_from_db()
        self.assertFalse(self.todo.done)
        self.assertIsNone(self.todo.completed_at)

    def test_an_agent_cannot_mark_a_colleagues_share(self):
        self.client.force_authenticate(user=self.ann)
        self.assertEqual(self.mark(self.bob.pk).status_code, 403)
        self.assertEqual(self.todo.assignee_completions.count(), 0)

    def test_an_admin_may_mark_anyones_share(self):
        self.client.force_authenticate(user=self.admin)
        self.assertEqual(self.mark(self.ann.pk).status_code, 200)
        completion = self.todo.assignee_completions.get(user=self.ann)
        self.assertEqual(completion.marked_by, self.admin, 'who pressed it is recorded')

    def test_someone_who_is_not_assigned_has_no_share_to_mark(self):
        outsider = User.objects.create_user(
            username='cout', full_name='Out', password='testpass123', role=User.Role.ADMIN
        )
        self.client.force_authenticate(user=self.admin)
        self.assertEqual(self.mark(outsider.pk).status_code, 404)

    def test_closing_the_todo_outright_finishes_every_share(self):
        self.client.force_authenticate(user=self.admin)
        response = self.client.patch(
            f'/api/todos/{self.todo.pk}/', {'done': True}, format='json'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.todo.assignee_completions.count(), 2)

    def test_reopening_the_todo_hands_every_share_back(self):
        self.client.force_authenticate(user=self.admin)
        self.client.patch(f'/api/todos/{self.todo.pk}/', {'done': True}, format='json')
        self.client.patch(f'/api/todos/{self.todo.pk}/', {'done': False}, format='json')
        self.assertEqual(self.todo.assignee_completions.count(), 0)

    def test_dropping_an_assignee_drops_their_tick(self):
        self.client.force_authenticate(user=self.ann)
        self.mark(self.ann.pk)
        self.client.force_authenticate(user=self.admin)
        self.client.patch(
            f'/api/todos/{self.todo.pk}/', {'assignee_ids': [self.bob.pk]}, format='json'
        )
        self.assertEqual(self.todo.assignee_completions.count(), 0)

    def test_someone_added_to_a_finished_todo_starts_finished(self):
        """Otherwise a done item would show an outstanding share, which it does not have."""
        self.client.force_authenticate(user=self.admin)
        self.client.patch(f'/api/todos/{self.todo.pk}/', {'done': True}, format='json')
        cid = User.objects.create_user(
            username='ccid', full_name='Cid', password='testpass123', role=User.Role.AGENT
        ).pk
        self.client.patch(
            f'/api/todos/{self.todo.pk}/',
            {'assignee_ids': [self.ann.pk, self.bob.pk, cid]},
            format='json',
        )
        self.todo.refresh_from_db()
        self.assertTrue(self.todo.done)
        self.assertEqual(self.todo.assignee_completions.count(), 3)


class PrivateTodoIsUnassignableTests(TestCase):
    """A private to-do is its author's alone.

    Nobody else can read one, so assigning it would hand somebody work they cannot open.
    The rule lives in the serializer rather than the UI, so it holds for the API too.
    """

    def setUp(self):
        self.client = APIClient()
        self.owner = User.objects.create_user(
            username='powner', full_name='Owner', password='testpass123', role=User.Role.AGENT
        )
        self.mate = User.objects.create_user(
            username='pmate', full_name='Mate', password='testpass123', role=User.Role.AGENT
        )
        self.client.force_authenticate(user=self.owner)

    def create(self, **extra):
        return self.client.post('/api/todos/', {'title': 'A job', **extra}, format='json')

    def test_a_private_todo_cannot_be_created_with_assignees(self):
        response = self.create(is_private=True, assignee_ids=[self.mate.pk])
        self.assertEqual(response.status_code, 400)
        self.assertIn('assignee_ids', response.json())

    def test_a_private_todo_with_nobody_on_it_is_fine(self):
        self.assertEqual(self.create(is_private=True).status_code, 201)

    def test_a_shared_todo_can_still_be_assigned(self):
        response = self.create(is_private=False, assignee_ids=[self.mate.pk])
        self.assertEqual(response.status_code, 201)

    def test_an_assigned_todo_cannot_be_flipped_to_private(self):
        """The assignees have to go first — flipping would strand them on unreadable work."""
        created = self.create(assignee_ids=[self.mate.pk]).json()
        response = self.client.patch(
            f'/api/todos/{created["id"]}/', {'is_private': True}, format='json'
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('assignee_ids', response.json())

    def test_clearing_the_assignees_in_the_same_request_is_accepted(self):
        """What the form actually sends: the whole payload, with nobody on it."""
        created = self.create(assignee_ids=[self.mate.pk]).json()
        response = self.client.patch(
            f'/api/todos/{created["id"]}/',
            {'is_private': True, 'assignee_ids': []},
            format='json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['is_private'])
        self.assertEqual(response.json()['assignees'], [])

    def test_a_private_todo_cannot_be_assigned_later(self):
        created = self.create(is_private=True).json()
        response = self.client.patch(
            f'/api/todos/{created["id"]}/', {'assignee_ids': [self.mate.pk]}, format='json'
        )
        self.assertEqual(response.status_code, 400)

    def test_a_private_folder_makes_the_rule_apply_too(self):
        """The folder decides privacy, so it can be what makes an item unassignable."""
        folder = self.client.post(
            '/api/todo-folders/', {'name': 'Mine', 'is_private': True}, format='json'
        ).json()
        response = self.create(folder_id=folder['id'], assignee_ids=[self.mate.pk])
        self.assertEqual(response.status_code, 400)
        self.assertIn('assignee_ids', response.json())
